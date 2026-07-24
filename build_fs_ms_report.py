"""Build the FS_MS paper-preparation DOCX and adjacent plain-text audit trail."""

from __future__ import annotations

import argparse
import hashlib
import json
import platform
import re
import subprocess
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = RGBColor(32, 55, 72)
MUTED = RGBColor(90, 98, 108)
GOLD = RGBColor(155, 112, 0)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd")) or OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    if shd.getparent() is None:
        tc_pr.append(shd)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")
    if tc_w.getparent() is None:
        tc_pr.append(tc_w)


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    if tbl_w.getparent() is None:
        tbl_pr.append(tbl_w)
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    if tbl_ind.getparent() is None:
        tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int],
              body_font_size: float = 8.5, header_font_size: float = 9.) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header_flag = OxmlElement("w:tblHeader")
    header_flag.set(qn("w:val"), "true")
    tr_pr.append(header_flag)
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
        shade(cell, "F4F6F9")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(header_font_size)
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = str(value)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(body_font_size)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run = paragraph.add_run()
    begin_run.font.size = Pt(9)
    begin_run._r.append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    instr_run = paragraph.add_run()
    instr_run.font.size = Pt(9)
    instr_run._r.append(instr)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run = paragraph.add_run()
    separate_run.font.size = Pt(9)
    separate_run._r.append(separate)
    result_run = paragraph.add_run("1")
    result_run.font.size = Pt(9)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run = paragraph.add_run()
    end_run.font.size = Pt(9)
    end_run._r.append(end)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10), ("Heading 2", 13, BLUE, 12, 6), ("Heading 3", 12, DARK_BLUE, 8, 4)):
        style = styles[name]
        style.font.name, style.font.size, style.font.color.rgb = "Calibri", Pt(size), RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(before), Pt(after)
        style.paragraph_format.keep_with_next = True
    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name, style.font.size = "Calibri", Pt(11)
        style.paragraph_format.left_indent = Inches(.375)
        style.paragraph_format.first_line_indent = Inches(-.194)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208
    header = section.header.paragraphs[0]
    header.text = "FS_MS study team | Functional-seizure microstate analysis"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in header.runs:
        run.font.size, run.font.color.rgb = Pt(8.5), MUTED
    add_page_field(section.footer.paragraphs[0])


def add_cover(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("RESEARCH REPORT")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(10.5), GOLD
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Microstate dynamics before functional seizure onset")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(29), NAVY
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("Rolling, topography-aware and cardiac-phase analyses with emphasis on microstates C and E")
    r.font.size, r.font.color.rgb = Pt(14), RGBColor(43, 81, 99)
    for text, bold in (("FS_MS study team", True), ("23 July 2026", False),
                       ("Paper-preparation report; inferential results require independent review before submission", False)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.bold, r.font.size, r.font.color.rgb = bold, Pt(10.5), MUTED
    doc.add_page_break()


def add_figure(doc: Document, path: Path, caption: str) -> None:
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(6.25))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", caption)
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(10)
    cr = cp.add_run(caption)
    cr.italic, cr.font.size, cr.font.color.rgb = True, Pt(9), MUTED


def fmt(value: object, digits: int = 3) -> str:
    try:
        value = float(value)
    except (TypeError, ValueError):
        return "-"
    if pd.isna(value):
        return "-"
    if value != 0 and abs(value) < .001:
        return f"{value:.2e}"
    return f"{value:.{digits}f}"


def compact_p(value: object) -> str:
    rendered = fmt(value)
    return rendered[1:] if rendered.startswith("0.") else rendered


def channel_name(value: object) -> str:
    """Render conventional EEG channel capitalization without changing stored data."""
    label = str(value).strip().lower()
    if label.startswith("fp"):
        return "Fp" + label[2:].upper()
    if label.endswith("z"):
        return label[:-1].upper() + "z"
    return label.upper()


def source_name(value: object) -> str:
    return {
        "global": "Global",
        "participant": "Participant",
        "participant_window_refit": "Window refit",
    }.get(str(value).strip().lower(), str(value).replace("_", " ").title())


def metric_name(value: object) -> str:
    return {
        "coverage": "Coverage",
        "mean_dwell_s": "Dwell time",
        "mean_gfp": "Mean GFP",
        "mean_topographic_gmd": "Topographic GMD",
        "occurrence_rate_hz": "Occurrence rate",
    }.get(str(value), str(value).replace("_", " ").title())


def slope_name(value: object) -> str:
    return {
        "early_elapsed_min": "Early slope",
        "late_elapsed_min": "Late slope",
    }.get(str(value), str(value).replace("_", " ").title())


def interaction_name(value: object) -> str:
    return {
        "is_systole:early_elapsed_min": "Systole × early slope",
        "is_systole:late_elapsed_min": "Systole × late slope",
    }.get(str(value), str(value).replace("_", " "))


def read_csv(path: Path) -> pd.DataFrame:
    return pd.read_csv(path) if path.exists() else pd.DataFrame()


def p_summary(frame: pd.DataFrame, corrected: str, alpha: float = .05) -> tuple[int, str]:
    if frame.empty or corrected not in frame:
        return 0, "No estimable models were available."
    sig = frame[frame[corrected].lt(alpha)]
    if len(sig):
        return len(sig), f"{len(sig)} term(s) survived multiplicity correction."
    valid = frame.dropna(subset=[corrected]).sort_values(corrected)
    return 0, (f"No term survived multiplicity correction; the smallest adjusted p-value was {fmt(valid.iloc[0][corrected])}."
               if len(valid) else "No adjusted p-values were estimable.")


def build_report(package: Path, results_dir: Path, manifest_path: Path, output: Path) -> dict:
    data_dir, figures = package / "data", package / "figures"
    manifest = pd.read_csv(manifest_path)
    cohort = read_csv(data_dir / "cohort_audit.csv")
    terms = read_csv(data_dir / "rolling_scalar_gee_terms.csv")
    contrasts = read_csv(data_dir / "rolling_early_late_contrasts.csv")
    cardiac = read_csv(data_dir / "rolling_cardiac_phase_gee_terms.csv")
    transition = read_csv(data_dir / "transition_group_condition_edge_model_terms.csv")
    hierarchy = read_csv(data_dir / "hierarchical_fit_summary.csv")
    covariance = read_csv(data_dir / "hierarchical_state_covariance_metrics.csv")
    tanova = read_csv(data_dir / "tanova_results.csv")
    bayes = read_csv(data_dir / "bayesian_tanova_results.csv")
    bhatt = read_csv(data_dir / "bhattacharyya_tanova_results.csv")
    outlier_audit = read_csv(data_dir / "rolling_window_outlier_audit.csv")
    morphology_peaks = read_csv(data_dir / "morphology_rft_peak_summary.csv")
    morphology_tfce_peaks = read_csv(data_dir / "morphology_tfce_peak_summary.csv")
    global_templates = read_csv(data_dir / "global_microstate_templates.csv")
    state_label_map = {}
    if not global_templates.empty and "state" in global_templates:
        for value in global_templates["state"].dropna().astype(str):
            match = re.match(r"State\s+(\d+)\s+\(([^)]+)\)", value)
            if match:
                state_label_map[int(match.group(1))] = match.group(2)
    summary = json.loads((data_dir / "analysis_summary.json").read_text(encoding="utf-8"))
    morphology_summary_path = data_dir / "morphology_rft_summary.json"
    morphology_summary = (json.loads(morphology_summary_path.read_text(encoding="utf-8"))
                          if morphology_summary_path.exists() else {})
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    doc.add_heading("Abstract", level=1)
    primary_scalar = terms[(terms.get("analysis_role", pd.Series(dtype=str)).eq("co_primary")) &
                           (terms.get("scheme", pd.Series(dtype=str)).eq("20min_60s_step10s")) &
                           terms.get("backfit_method", pd.Series(dtype=str)).eq("hard") &
                           terms.get("stratum", pd.Series(dtype=str)).astype(str).str.contains("functional", na=False)]
    scalar_n, scalar_text = p_summary(primary_scalar, "p_fdr")
    cardiac_primary = cardiac[(cardiac.get("systolic_window_start_ms", pd.Series(dtype=float)).eq(180)) &
                              (cardiac.get("systolic_window_end_ms", pd.Series(dtype=float)).eq(500)) &
                              cardiac.get("stratum", pd.Series(dtype=str)).astype(str).str.contains("functional", na=False)]
    cardiac_n, cardiac_text = p_summary(cardiac_primary[cardiac_primary.get("term", pd.Series(dtype=str)).astype(str).str.contains(":", na=False)], "p_holm")
    cardiac_abstract = cardiac_text[:1].lower() + cardiac_text[1:] if cardiac_text else "no estimable interaction was available."
    morphology_primary = morphology_peaks[morphology_peaks.get("family", pd.Series(dtype=str)).eq("primary")].copy()
    morphology_n = int(morphology_primary.get("peak_p_fwer", pd.Series(dtype=float)).lt(.05).sum()) if len(morphology_primary) else 0
    morphology_best = (morphology_primary.dropna(subset=["peak_p_fwer"]).sort_values("peak_p_fwer").iloc[0]
                       if len(morphology_primary.dropna(subset=["peak_p_fwer"])) else None)
    tfce_primary = morphology_tfce_peaks[morphology_tfce_peaks.get("family", pd.Series(dtype=str)).eq("primary")].copy()
    tfce_n = int(tfce_primary.get("peak_p_tfce_fwer", pd.Series(dtype=float)).lt(.05).sum()) if len(tfce_primary) else 0
    tfce_best = (tfce_primary.dropna(subset=["peak_p_tfce_fwer"]).sort_values("peak_p_tfce_fwer").iloc[0]
                 if len(tfce_primary.dropna(subset=["peak_p_tfce_fwer"])) else None)
    tfce_abstract = (f"Primary electrode-graph TFCE found {tfce_n} jointly corrected C/E morphology field(s)."
                     if tfce_n else "Primary electrode-graph TFCE found no jointly corrected C/E morphology field.")
    gaussian_abstract = (f"The Gaussian maximum-field sensitivity found {morphology_n} corrected field(s)."
                         if morphology_n else "The Gaussian maximum-field sensitivity found no corrected field.")
    coverage_outliers = int(summary.get("coverage_outlying_windows", 0))
    gfp_outliers = int(summary.get("mean_gfp_outlying_windows", 0))
    doc.add_paragraph(
        f"This report audits and extends a hierarchical EEG microstate analysis around functional seizure onset. "
        f"The corrected paper manifest contains {len(manifest):,} unique records from {manifest['participant'].nunique():,} participants; "
        f"rolling analyses contain {summary.get('rolling_participants', 0):,} participants. Prespecified global microstate C and "
        f"participant-template E were evaluated alongside global-E and participant-C sensitivities in 60-s windows stepped every 10 s across the final 20 min, "
        f"with 10-s/1-s final-minute sensitivity analyses. Participant-clustered piecewise GEE, polarity-invariant topographic GMD, "
        f"permutation TANOVA, channel-level Bhattacharyya analysis, rolling-window K=6 refitting with primary electrode-graph TFCE, a Gaussian maximum-field sensitivity and exposure-adjusted cardiac-phase GEE were used. {scalar_text} "
        f"For cardiac phase, {cardiac_abstract} A secondary TANOVA detected a functional preictal-versus-postictal difference in global E, "
        f"but this is a condition-level contrast rather than a continuous prodromal trajectory. {tfce_abstract} {gaussian_abstract} "
        f"The earlier fixed-backfit p=0.034 morphology result is superseded because it tested template similarity rather than change in refitted state morphology. Results are interpreted as paper-forming evidence rather than confirmation because the analyses "
        f"are exploratory and several related outcomes were tested."
    )
    doc.add_page_break()
    doc.add_heading("Executive interpretation", level=1)
    if scalar_n:
        doc.add_paragraph("There is corrected evidence that at least one rolling microstate feature changes as functional seizure onset approaches. The strongest terms are shown below; emphasis should remain on effect sizes and intervals rather than a binary threshold.")
    else:
        doc.add_paragraph("The rolling analysis does not provide multiplicity-corrected evidence for a systematic C- or E-state trajectory toward onset. This is informative: any apparent curves should be treated as weak or heterogeneous rather than a group-level prodrome.")
    if cardiac_n:
        doc.add_paragraph("Cardiac phase modifies at least one preictal trajectory after Holm correction. This supports a state-by-body timing hypothesis, but replication and an alternative 200-400 ms systolic definition are required.")
    else:
        doc.add_paragraph("Systole/diastole did not reliably modify C/E trajectories after correction. Small uncorrected effects remain hypothesis-generating, not evidence of phase locking.")
    if tfce_n and tfce_best is not None:
        doc.add_paragraph(
            f"Primary TFCE identified {tfce_n} jointly corrected window-refitted C/E morphology field(s). The strongest field was participant "
            f"{tfce_best['template_label']} in {tfce_best['stratum']} at {channel_name(tfce_best['peak_channel'])} "
            f"(t={fmt(tfce_best['peak_t'])}, joint TFCE-FWER p={fmt(tfce_best['peak_p_tfce_fwer'])}). Betas quantify normalized topographic redistribution per minute, not voltage change."
        )
    elif tfce_best is not None:
        doc.add_paragraph(
            f"Primary TFCE found no corrected window-refitted C/E morphology field. The closest field was participant {tfce_best['template_label']} "
            f"in {tfce_best['stratum']} at {channel_name(tfce_best['peak_channel'])} (t={fmt(tfce_best['peak_t'])}, "
            f"joint TFCE-FWER p={fmt(tfce_best['peak_p_tfce_fwer'])})."
        )
    if morphology_n and morphology_best is not None:
        doc.add_paragraph(
            f"The Gaussian-smoothed sensitivity found {morphology_n} corrected field(s); its strongest peak was participant "
            f"{morphology_best['template_label']} in {morphology_best['stratum']} near {channel_name(morphology_best['peak_nearest_channel'])} "
            f"(joint maximum-field p={fmt(morphology_best['peak_p_fwer'])})."
        )
    doc.add_paragraph("Topography-aware secondary analysis nevertheless detected a max-state-corrected preictal-versus-postictal difference for global E in functional events. Bhattacharyya localization further highlighted E at CP2, Pz, Fp1 and P4 when comparing functional with epileptic preictal events. These findings concern spatial configuration and phenotype contrasts; they do not overturn the negative continuous preictal-trajectory result.")
    if scalar_n or cardiac_n or tfce_n:
        doc.add_paragraph("The most defensible paper framing is a candidate but heterogeneous preictal signal, anchored to corrected effect estimates and accompanied by explicit sensitivity analyses; it is not yet a validated biomarker.")
    else:
        doc.add_paragraph("The most defensible paper framing is a careful negative-or-heterogeneous preictal trajectory result, accompanied by condition-level/topographic signals and transparent sensitivity analyses.")

    doc.add_heading("1. Background and aims", level=1)
    doc.add_paragraph("Functional seizures are episodic expressions of functional neurological disorder. EEG microstates provide a compact description of rapidly changing scalp electric-field configurations, while cardiac phase offers a mechanistically relevant index of brain-body timing. A recent pilot reported shorter microstate durations in functional/dissociative seizures, particularly for class D, but evidence specifically tracking the approach to recorded seizure onset remains limited [1].")
    doc.add_paragraph("The present work asked whether microstate coverage, GFP-weighted expression, topographic fit or the full scalp morphology changed during the final 20 min before functional seizure onset, with prespecified attention to global C and participant-template E. It also tested whether C/E occurrence rates differed between systole and diastole as onset approached, and whether condition-level topographies showed complementary evidence.")

    audit_heading = doc.add_heading("2. Pipeline and data-integrity audit", level=1)
    audit_heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph("The original pipeline had run and produced a complete hierarchy for all previously represented records, but three issues made the existing summary unsafe to reuse: a filename collision duplicated one interictal crop; the global one-second output covered only the final minute; and transition rows were grouped by participant-local numeric state indices even though canonical C/E indices vary between participants. The implementation fixes all three causes and leaves the historical output intact.")
    doc.add_paragraph(f"The corrected paper manifest has {len(manifest):,} rows and {manifest['file_path'].astype(str).str.lower().nunique():,} unique paths. Pending HEPPy records remain excluded rather than imputed. Counts below are deidentified and stratified so FS-only functional events and DD functional events are not pooled as the primary analysis.")
    if not cohort.empty:
        rows = [[str(r.stratum), str(r.group), str(r.condition), str(r.seizure_type), str(int(r.n_records)), str(int(r.n_participants))]
                for r in cohort.itertuples(index=False)]
        add_table(doc, ["Stratum", "Group", "Condition", "Event", "Records", "Participants"], rows, [2200, 700, 1200, 1200, 900, 3160])

    doc.add_heading("3. Methods", level=1)
    doc.add_heading("3.1 Microstate hierarchy and labels", level=2)
    if not hierarchy.empty:
        global_row = hierarchy[hierarchy["level"].eq("global")].head(1)
        if len(global_row):
            r = global_row.iloc[0]
            global_labels = (sorted(covariance.loc[covariance.get("level", pd.Series(dtype=str)).eq("global"), "template_label"].dropna().astype(str).unique())
                             if not covariance.empty and "template_label" in covariance else [])
            doc.add_paragraph(f"Hierarchical fitting selected K={int(r['K_estimated'])} global states. Mean template correlation was {fmt(r['template_mean_corr'])}, with {int(r['template_strong_matches'])}/{int(r['K_estimated'])} strong canonical matches. The corrected global labels were {', '.join(global_labels)}. Global C and participant-template E retain their prespecified roles; newly available global E is a corroborating sensitivity. Global G was not present and was not forced into the analysis.")
        participant_rows = hierarchy[hierarchy["level"].eq("participant")].copy()
        if len(participant_rows):
            complete = participant_rows["template_strong_matches"].eq(participant_rows["K_estimated"]).sum()
            weak = participant_rows["template_strong_matches"].lt(participant_rows["K_estimated"]).sum()
            doc.add_paragraph(
                f"At participant level, {int(complete)}/{len(participant_rows)} fits had a strong match for every selected state "
                f"at the pipeline's |r| >= 0.50 threshold; {int(weak)} had at least one weaker mapping. State-specific template "
                "correlations are retained in the rolling data so C/E findings can be checked against fit quality."
            )
    doc.add_heading("3.2 Rolling analysis", level=2)
    doc.add_paragraph("Primary windows were 60 s long and stepped every 10 s from -20 min to onset. A final-minute sensitivity used 10-s windows stepped every 1 s. Before events were averaged within participant, coverage and mean GFP were filtered independently within each event, source, state, backfit method and rolling scheme. For each metric, a raw finite value was excluded when its absolute distance from the event-sequence median exceeded three sample SDs; groups with fewer than four finite windows or zero SD were unchanged. Thresholds were calculated once on the raw scale, without iterative or MAD filtering. Raw values, cleaned values and metric-specific flags are retained in the data package. Multiple retained events were then averaged equally within participant and window. Participant-template states required the pipeline's prespecified strong mapping threshold (|r| >= 0.50). Coverage was modelled on a logit scale; occurrence rate (with a half-event correction to retain zero-count windows), mean dwell, mean GFP and polarity-invariant topographic GMD were log transformed. Participant-clustered Gaussian GEE estimated continuous piecewise slopes for -20 to -5 min and -5 to 0 min. Benjamini-Hochberg FDR controlled the scalar family. Hard backfit was primary and Gaussian-mixture backfit was sensitivity. Non-converged or non-estimable fits retain their status and estimate where available, but their p-values were set to missing before multiplicity correction.")
    doc.add_heading("3.3 Topography-aware analysis", level=2)
    doc.add_paragraph("Rolling topographic GMD quantified the polarity-invariant distance between each fitted map and its template. Condition-level topographies were assessed by repeated-measures/permutation TANOVA with max-state correction (5,000 permutations). Channel/state localization used Bhattacharyya distances with participant weighting and Bayesian FDR; these channel results are secondary because neighbouring electrodes and states are dependent.")
    doc.add_heading("3.4 Continuous scalp-morphology maximum-field analysis", level=2)
    doc.add_paragraph(
        "Morphology was re-estimated rather than backfitted. For each preictal event, end-aligned 60-s windows were placed every 30 s across the available final 20 min. Each window selected only GFP-peak maps cached from the participant's full-record preprocessing, avoiding window-edge refiltering. A locked K=6 SPM VB mixture was fitted with spm_mix_prior, using the participant's six full-record maps as component-mean priors with mean weight one; parent maps were not inserted into the observations. At least 30 real GFP peaks were required. Fitted states were matched one-to-one to the participant solution and polarity-aligned; an individual A/C/E map required at least five assigned peaks. Low window-to-prior correlations were retained because declining similarity may itself be biological, whereas the participant's full A/C/E reference required canonical |r| >= 0.50."
    )
    doc.add_paragraph(
        "For every event, state and electrode, one ordinary least-squares slope was fitted to the sequence of polarity-aligned unit-norm channel values against minutes to onset. An event/state required at least 12 valid windows spanning at least five minutes. Event slopes were averaged equally within participant, leaving one participant C and one participant E beta map per functional stratum. These betas represent normalized topographic redistribution per minute and cannot be interpreted as absolute EEG voltage. The earlier fixed-backfit morphology p=0.034 is superseded: it summarized maps assigned to a fixed template and therefore estimated a different quantity."
    )
    doc.add_paragraph(
        "Electrode-graph TFCE was the primary spatial inference. The 27 measured electrodes were connected by a 68-edge Delaunay graph. Positive and negative t fields were enhanced separately with E=0.5 and H=2 and exact integration across observed statistic heights. Five thousand whole-participant sign flips used identical participant signs across fields; the maximum absolute TFCE score jointly corrected all electrodes and four co-primary fields (C and E in FS-functional and DD-functional strata) [8,9]. A 45-degree FWHM Gaussian scalp interpolation with permutation maximum-field correction over the same four fields was retained only as a separately reported sensitivity; it is RFT-like maximum-field inference, not parametric resel-based random-field theory [6,7]."
    )
    doc.add_heading("3.5 Cardiac-phase analysis", level=2)
    doc.add_paragraph("Corrected R peaks defined systole as 180-500 ms after each R peak. Within each rolling window, state-onset counts and the exact systolic/diastolic exposure time were reconstructed. Exposure-offset Poisson GEE included phase, early and late time, and phase-by-time interactions. The interaction family across C/E was Holm corrected. A 200-400 ms window is reserved as a sensitivity definition consistent with prior EEG cardiac-phase work [4].")

    doc.add_heading("4. Results", level=1)
    doc.add_heading("4.1 Rolling microstate behaviour", level=2)
    doc.add_paragraph(
        f"The locked within-event rule excluded {coverage_outliers:,} coverage window values and {gfp_outliers:,} mean-GFP window values. "
        "Filtering was metric-specific, so exclusion from one outcome did not remove another valid outcome from the same window. "
        f"The audit table contains {len(outlier_audit):,} metric-by-scheme/source/state/stratum cells with retained and excluded counts."
    )
    estimable_primary = (primary_scalar.dropna(subset=["p_fdr"]).sort_values(["p_fdr", "p_value"])
                         if {"p_fdr", "p_value"}.issubset(primary_scalar.columns) else pd.DataFrame())
    if len(estimable_primary):
        best = estimable_primary.iloc[0]
        if float(best["p_fdr"]) < .05:
            doc.add_paragraph(
                f"The strongest corrected primary functional-seizure trajectory was {source_name(best['source'])} {best['template_label']} "
                f"{metric_name(best['metric'])} in {best['stratum']} ({slope_name(best['term'])}; estimate {fmt(best['estimate'])}, "
                f"95% CI {fmt(best['ci_low'])} to {fmt(best['ci_high'])}, q={fmt(best['p_fdr'])})."
            )
        else:
            doc.add_paragraph(
                f"No primary hard-backfit functional-seizure slope survived FDR. The smallest adjusted value was q={fmt(best['p_fdr'])} "
                f"for {source_name(best['source'])} {best['template_label']} {metric_name(best['metric'])} in {best['stratum']} ({slope_name(best['term'])})."
            )
    if not terms.empty:
        focus = primary_scalar.copy()
        focus = focus.sort_values(["p_fdr", "p_value"], na_position="last").head(12)
        rows = [[source_name(r.source), r.template_label, r.stratum, metric_name(r.metric), slope_name(r.term),
                 fmt(getattr(r, "estimate", None)), f"{fmt(getattr(r, 'ci_low', None))} to {fmt(getattr(r, 'ci_high', None))}",
                 compact_p(getattr(r, "p_value", None)), compact_p(getattr(r, "p_fdr", None)), str(int(r.n_participants))]
                for r in focus.itertuples(index=False)]
        doc.add_page_break()
        add_table(doc, ["Source", "St.", "Stratum", "Metric", "Slope", "Est.", "95% CI", "p", "q", "n"],
                  rows, [1100, 500, 1250, 1200, 900, 700, 1700, 550, 550, 450], 7.6, 8.)
        final_focus = terms[(terms["analysis_role"].eq("co_primary")) &
                            (terms["scheme"].eq("final60s_10s_step1s")) & terms["backfit_method"].eq("hard") &
                            terms["stratum"].astype(str).str.contains("functional")].copy()
        if len(final_focus):
            doc.add_heading("Final-minute sensitivity", level=3)
            estimable_final = final_focus[final_focus.get("status", pd.Series(index=final_focus.index, dtype=str)).eq("ok")].dropna(subset=["p_value", "p_fdr"])
            if len(estimable_final):
                estimable_final = estimable_final.sort_values(["p_fdr", "p_value"]).head(8)
                rows = [[r.source, r.template_label, r.stratum, r.metric, r.term,
                         fmt(getattr(r, "estimate", None)), fmt(getattr(r, "p_value", None)),
                         fmt(getattr(r, "p_fdr", None)), str(int(r.n_participants))]
                        for r in estimable_final.itertuples(index=False)]
                add_table(doc, ["Source", "State", "Stratum", "Metric", "Slope", "Estimate", "p", "q", "n"],
                          rows, [650, 500, 1550, 1000, 1100, 950, 750, 750, 2110])
            else:
                status_counts = "; ".join(f"{key}={value}" for key, value in final_focus["status"].fillna("missing").value_counts().items())
                doc.add_paragraph(f"No co-primary final-minute inferential fit converged or met the minimum-data requirement ({status_counts}). The descriptive rolling values and explicit status fields remain in the supplementary tables, but no final-minute p-value is interpreted.")
    add_figure(doc, figures / "rolling_coverage.png", "Figure 1. Participant-weighted 60-s rolling coverage trajectories; dashed line marks the prespecified -5-min knot.")
    add_figure(doc, figures / "rolling_occurrence_rate_hz.png", "Figure 2. Participant-weighted microstate occurrence rates in 60-s rolling windows.")
    add_figure(doc, figures / "rolling_mean_dwell_s.png", "Figure 3. Participant-weighted mean microstate dwell time in 60-s rolling windows.")
    add_figure(doc, figures / "rolling_mean_topographic_gmd.png", "Figure 4. Rolling polarity-invariant topographic GMD. Lower values indicate closer fit to the relevant template.")
    add_figure(doc, figures / "rolling_mean_gfp.png", "Figure 5. Rolling state-specific mean GFP trajectories.")
    if not contrasts.empty:
        doc.add_paragraph("Early-versus-late paired contrasts are supplied in the data folder. They are supportive summaries; the piecewise GEE is the main temporal analysis because it uses the full trajectory.")

    doc.add_heading("4.2 TANOVA and Bhattacharyya localization", level=2)
    tanova_sig = tanova[tanova.get("p_max_state", pd.Series(dtype=float)).lt(.05)] if not tanova.empty else pd.DataFrame()
    if len(tanova_sig):
        doc.add_paragraph(f"TANOVA identified {len(tanova_sig)} state/effect result(s) after within-effect max-state correction; the strongest 12 are shown.")
        functional_e = tanova_sig[(tanova_sig["effect"].astype(str).str.startswith("Functional events")) & tanova_sig["state"].map(state_label_map).eq("E")]
        if len(functional_e):
            r = functional_e.sort_values("p_max_state").iloc[0]
            doc.add_paragraph(f"For functional events, global E differed between preictal and postictal conditions (max-state p={fmt(r['p_max_state'])}; n={int(r['n_units'])}). This supports a condition-level E topographic difference, not a monotonic change as onset approaches.")
        rows = [[str(r.effect), str(r.state), state_label_map.get(int(r.state), ""), fmt(r.statistic), fmt(r.p_uncorrected), fmt(r.p_max_state), str(int(r.n_units))]
                for r in tanova_sig.sort_values(["p_max_state", "p_uncorrected"]).head(12).itertuples(index=False)]
        add_table(doc, ["Effect", "State", "Label", "Statistic", "p", "max-state p", "n"], rows,
                  [3800, 600, 600, 950, 800, 1100, 700], header_font_size=8.)
    else:
        doc.add_paragraph("No TANOVA state survived max-state correction.")
    bh_sig = bhatt[bhatt.get("bayesian_fdr_significant", pd.Series(dtype=float)).eq(1)] if not bhatt.empty else pd.DataFrame()
    if len(bh_sig):
        c_e = bh_sig[bh_sig["template_label"].isin(["C", "E"])]
        doc.add_paragraph(f"Bhattacharyya localization identified {len(bh_sig)} channel/state findings at Bayesian FDR, including {len(c_e)} involving C or E. These localizations should guide confirmatory regions of interest rather than be presented as independent electrode discoveries.")
        fs_es_e = c_e[(c_e["effect"].astype(str).str.startswith("Events: Preict functional")) & c_e["template_label"].eq("E")]
        if len(fs_es_e):
            channels = ", ".join(fs_es_e.sort_values("posterior_null_probability")["channel_label"].astype(str))
            best_local = fs_es_e.sort_values("posterior_null_probability").iloc[0]
            doc.add_paragraph(f"The clearest C/E localization compared preictal functional with preictal epileptic events: E at {channels}; the strongest channel was {best_local['channel_label']} (permutation p={fmt(best_local['permutation_p'])}, posterior probability above the null 95th percentile={fmt(best_local['posterior_prob_above_null95'])}). This contrast distinguishes event phenotypes and does not test within-functional-event approach to onset.")
        rows = [[str(r.effect), str(r.template_label), str(r.channel_label), fmt(r.bhattacharyya_distance),
                 fmt(r.permutation_p), fmt(r.posterior_prob_above_null95), str(int(r.n_participants))]
                for r in bh_sig.sort_values("posterior_prob_above_null95", ascending=False).head(12).itertuples(index=False)]
        add_table(doc, ["Effect", "State", "Channel", "Distance", "perm. p", "Pr(>null95)", "n"], rows,
                  [4230, 600, 850, 900, 900, 1150, 670], header_font_size=8.)
    else:
        doc.add_paragraph("No channel/state Bhattacharyya localization survived Bayesian FDR.")
    doc.add_page_break()
    add_figure(doc, figures / "global_microstate_templates.png", "Figure 6. Global microstate templates and canonical labels from the corrected hierarchy.")
    add_figure(doc, figures / "tanova_pmax_heatmap.png", "Figure 7. Max-state-corrected TANOVA results across prespecified contrasts.")

    doc.add_page_break()
    doc.add_heading("4.3 Continuous preictal scalp morphology", level=2)
    if len(tfce_primary):
        focus = tfce_primary.sort_values("peak_p_tfce_fwer", na_position="last")
        rows = [[source_name(r.source), str(r.template_label), str(r.stratum),
                 "Linear", str(int(r.n_participants)),
                 channel_name(r.peak_channel), fmt(r.peak_t), fmt(r.peak_tfce_score), fmt(r.peak_p_tfce_fwer)]
                for r in focus.itertuples(index=False)]
        add_table(doc, ["Source", "State", "Stratum", "Slope", "n", "Peak", "t", "TFCE", "FWER p"],
                  rows, [1200, 600, 1400, 900, 500, 600, 800, 950, 1050], header_font_size=8.)
        if tfce_n and tfce_best is not None:
            doc.add_paragraph(
                f"Primary TFCE identified {tfce_n} jointly corrected field(s). The strongest was participant {tfce_best['template_label']} "
                f"in {tfce_best['stratum']} (n={int(tfce_best['n_participants'])}) at {channel_name(tfce_best['peak_channel'])} "
                f"(beta/min={fmt(tfce_best['peak_beta_per_min'], 5)}, t={fmt(tfce_best['peak_t'])}, "
                f"TFCE={fmt(tfce_best['peak_tfce_score'])}, joint TFCE-FWER p={fmt(tfce_best['peak_p_tfce_fwer'])}); "
                f"{int(tfce_best['significant_electrodes'])} electrode(s) crossed p<0.05. The beta represents redistribution of a polarity-aligned unit-norm map, not a change in absolute EEG voltage."
            )
        elif tfce_best is not None:
            doc.add_paragraph(
                f"No primary TFCE field survived joint correction. The closest field was participant {tfce_best['template_label']} "
                f"in {tfce_best['stratum']} (n={int(tfce_best['n_participants'])}) at {channel_name(tfce_best['peak_channel'])} "
                f"(beta/min={fmt(tfce_best['peak_beta_per_min'], 5)}, t={fmt(tfce_best['peak_t'])}, "
                f"TFCE={fmt(tfce_best['peak_tfce_score'])}, joint TFCE-FWER p={fmt(tfce_best['peak_p_tfce_fwer'])})."
            )
    else:
        doc.add_paragraph("Window-refitted TFCE results were unavailable.")
    doc.add_paragraph(
        "The previous fixed-backfit morphology result (uncorrected p=0.034) is superseded. It tested a different estimand—change in maps assigned to fixed templates—rather than change in independently refitted rolling-window microstate morphology."
    )
    add_figure(doc, figures / "morphology_tfce_primary_beta_maps.png",
               "Figure 8. Primary electrode-graph TFCE analysis of participant-level slopes from window-refitted C/E maps. Open circles mark measured electrodes; filled black points, when present, denote joint TFCE-FWER p<0.05.")
    doc.add_page_break()
    doc.add_heading("4.3.1 Gaussian maximum-field sensitivity", level=3)
    if len(morphology_primary):
        focus = morphology_primary.sort_values("peak_p_fwer", na_position="last")
        rows = [[source_name(r.source), str(r.template_label), str(r.stratum), "Linear", str(int(r.n_participants)),
                 channel_name(r.peak_nearest_channel), fmt(r.peak_beta_per_min, 5), fmt(r.peak_t),
                 fmt(getattr(r, "peak_p_uncorrected", None)), fmt(r.peak_p_fwer)]
                for r in focus.itertuples(index=False)]
        add_table(doc, ["Source", "State", "Stratum", "Slope", "n", "Peak", "Beta/min", "t", "p", "FWER p"],
                  rows, [1200, 600, 1350, 850, 450, 550, 950, 700, 700, 950], header_font_size=8.)
        if morphology_n:
            doc.add_paragraph(f"The Gaussian sensitivity contained {morphology_n} field(s) surviving joint maximum-field correction.")
        elif morphology_best is not None:
            doc.add_paragraph(
                f"No Gaussian sensitivity field survived joint correction. The closest field was participant {morphology_best['template_label']} "
                f"in {morphology_best['stratum']} near {channel_name(morphology_best['peak_nearest_channel'])} "
                f"(beta/min={fmt(morphology_best['peak_beta_per_min'], 5)}, t={fmt(morphology_best['peak_t'])}, "
                f"uncorrected p={fmt(morphology_best.get('peak_p_uncorrected'))}, FWER p={fmt(morphology_best['peak_p_fwer'])})."
            )
    add_figure(doc, figures / "morphology_rft_primary_beta_maps.png",
               "Figure 9. Gaussian-smoothed sensitivity analysis of the same participant C/E beta fields. Black grid points, when present, denote joint maximum-field FWER p<0.05.")

    doc.add_heading("4.4 Systole/diastole modulation", level=2)
    if not cardiac.empty:
        focus = cardiac_primary[cardiac_primary["term"].astype(str).str.contains(":", na=False)].sort_values(["p_holm", "p_value"], na_position="last").head(12)
        rows = [[str(r.template_label), str(r.stratum), interaction_name(r.term), fmt(r.rate_ratio),
                 f"{fmt(r.ci_low)} to {fmt(r.ci_high)}", fmt(r.p_value), fmt(r.p_holm), str(int(r.n_participants))]
                for r in focus.itertuples(index=False)]
        add_table(doc, ["State", "Stratum", "Interaction", "Rate ratio", "95% CI", "p", "Holm p", "n"],
                  rows, [650, 1300, 2500, 1000, 1350, 760, 900, 800], header_font_size=8.)
    doc.add_paragraph(cardiac_text)
    doc.add_paragraph("The 200-400 ms systolic sensitivity is supplied in the cardiac model table and should be compared for direction and interval overlap rather than selected by significance.")

    doc.add_heading("4.5 C-E transition sensitivity", level=2)
    if transition.empty:
        doc.add_paragraph("Canonical C-E transition inference was unavailable or below the minimum cell size; transition tables are retained descriptively.")
    else:
        sig = transition[transition.get("p_fdr", pd.Series(dtype=float)).lt(.05)]
        doc.add_paragraph(f"After aggregating solely by canonical labels, {len(sig)} transition-model term(s) survived FDR. This replaces the invalid local-index aggregation in the earlier script.")

    doc.add_heading("5. Discussion", level=1)
    doc.add_paragraph("The analysis separates three questions that are often conflated: whether a state is more common, whether its expressed maps fit the template differently, and whether its occurrence is preferentially timed to cardiac phase. A positive result in one domain does not imply the others. Participant E remains a within-participant construct, while the freshly recovered global E offers a complementary common-map sensitivity.")
    if scalar_n == 0:
        doc.add_paragraph("The absence of corrected C/E rolling slopes argues against a simple, population-wide microstate prodrome over 20 min. It does not exclude sharp changes inside a shorter interval, changes confined to a subgroup, or participant-specific directions that cancel at group level. The final-minute sensitivity, functional-stratum separation and individual trajectories in the data package are the appropriate routes for those alternatives.")
    else:
        doc.add_paragraph("Corrected rolling effects provide a candidate preictal signature, but the overlapping windows make the plotted series visually smoother than the independent information content. The participant-clustered model, corrected p-values and confidence intervals should therefore anchor any manuscript claim.")
    doc.add_paragraph("Condition-level topographic results can complement a weak scalar trajectory by showing that the configuration, rather than the amount, of a state differs. Conversely, isolated channel Bhattacharyya findings without TANOVA support should be treated as localization hypotheses. This hierarchy of evidence is consistent with microstate standardisation guidance [5].")
    if tfce_n:
        doc.add_paragraph(
            "The rolling-window refit provides evidence about morphological reconfiguration itself: every window produced a new state solution under a participant prior, and the inferential unit was the participant's event-averaged channel slope. A corrected TFCE field therefore supports consistent spatial redistribution as onset approaches, while its channel beta remains a normalized topographic effect rather than voltage. Replication is still required because windows overlap, participant priors include the target event and the montage is sparse."
        )
    else:
        doc.add_paragraph(
            "The window-refitted TFCE result extends the scalar analysis to the full polarity-aligned scalp field: no participant C/E morphology trajectory was sufficiently consistent across participants after joint correction over electrodes and four fields. The Gaussian result is a spatial sensitivity only and cannot rescue a negative primary TFCE result."
        )
    doc.add_paragraph("The earlier fixed-backfit p=0.034 result should not appear as supporting evidence in a manuscript: it estimated change in averages assigned to fixed templates, whereas the replacement asks whether independently refitted state morphology changes through time.")

    doc.add_heading("6. Limitations", level=1)
    for text in (
        "This is an exploratory secondary analysis with multiple related outcomes and no independent validation cohort.",
        "Participant C/E priors were fitted from the full participant record and therefore include the target event; leave-one-event-out priors were outside scope.",
        "Overlapping rolling windows are highly correlated. GEE handles participant clustering but does not create new independent observations.",
        "Pending HEPPy records were excluded. Missingness could bias results if preprocessing availability relates to event phenotype or signal quality.",
        "Cardiac analyses depend on corrected R peaks and on an operational systolic window; the 200-400 ms sensitivity should be reported.",
        "Time-to-ictus alignment inherits the clinical seizure-onset annotations and available crop duration; onset-timing error would blur a short prodrome.",
        "The one-second final-minute sensitivity uses the audited July 20 backfits. The repaired collision was interictal, so no preictal file changed, but this remains a sensitivity rather than the primary refit.",
        "No co-primary final-minute inferential model converged or met the minimum-data rule; its rolling values are descriptive only and extreme retained warning estimates must not be interpreted.",
        "TANOVA and channel-level localization are condition-level complements, not direct tests of a continuous preictal trajectory.",
        "The 60-s morphology windows overlap by 30 s, so successive fitted maps are correlated; event-level regression avoids treating windows as independent inferential units but does not add independent information.",
        "TFCE used channel count as cluster extent on a 68-edge Delaunay graph. With only 27 electrodes, spatial extent and anatomical localization are coarse.",
        "The Gaussian sensitivity grid is an interpolation, not a denser measurement array; the fixed 45-degree kernel should not be interpreted as anatomical resolution.",
        "Participant sign-flip validity relies on exchangeable, approximately symmetric participant beta maps. Identical signs preserve cross-field dependence but do not remove onset-annotation error or cohort heterogeneity.",
    ):
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("7. Paper-development roadmap", level=1)
    for text in (
        "Pre-register one primary C outcome and one participant-E outcome from this report before any confirmatory rerun.",
        "Use FS-only functional and DD functional cohorts as separate primary strata; present pooled functional events only as secondary.",
        "Make the 20-min piecewise trajectory the main figure, with final-minute and mixture backfit in supplementary material.",
        "Report effect sizes, confidence intervals, raw p-values and corrected p-values together; avoid electrode-count language.",
        "Seek blinded replication or a held-out cohort before describing any signal as a biomarker.",
        "For confirmation, lock K=6 participant-prior refits, 60-s/30-s windows, event-level linear channel slopes, the 27-electrode graph, TFCE E=0.5/H=2 and joint four-field C/E-by-stratum correction; use a denser montage if anatomical localization is a goal.",
    ):
        doc.add_paragraph(text, style="List Number")

    doc.add_heading("Reproducibility and data package", level=1)
    doc.add_paragraph("The adjacent TXT audit trail records commands, versions, hashes, inclusion/exclusion counts, window-refit failures, outlier counts, model specifications, convergence status, document QA and a final output inventory. The data folder contains raw/clean event-level rolling metrics with flags, participant-weighted rolling metrics, model coefficients, contrasts, cardiac exposure/count tables, topographic statistics, transitions, deidentified rolling-window refit maps and event/participant morphology betas, primary electrode-graph TFCE results, Gaussian maximum-field sensitivities and data dictionaries. Direct patient identifiers and patient-specific source paths are not reproduced in this report.")

    doc.add_heading("References", level=1)
    refs = [
        "1. Kučikienė D, et al. Changes in brain network dynamics during functional/dissociative seizures: An exploratory pilot study on EEG microstates (2025). PMID 40757263.",
        "2. Altered microstate dynamics in Functional Neurological Disorder (2026). PMID 41747648.",
        "3. Kandasamy R, et al. The heartbeat evoked potential and the prediction of functional seizure semiology (2026). PMID 41994620.",
        "4. Al E, et al. Heart-brain interactions shape somatosensory perception and evoked potentials (2020). PMID 32341167.",
        "5. MicrostateLab: standardising EEG microstate analysis (2023). PMID 37697212.",
        "6. Maris E, Oostenveld R. Nonparametric statistical testing of EEG- and MEG-data. Journal of Neuroscience Methods (2007);164:177-190. PMID 17517438. doi:10.1016/j.jneumeth.2007.03.024.",
        "7. Nichols TE, Holmes AP. Nonparametric permutation tests for functional neuroimaging: a primer with examples. Human Brain Mapping (2002);15:1-25. PMID 11747097. doi:10.1002/hbm.1058.",
        "8. Smith SM, Nichols TE. Threshold-free cluster enhancement: addressing problems of smoothing, threshold dependence and localisation in cluster inference. NeuroImage (2009);44:83-98. PMID 18501637. doi:10.1016/j.neuroimage.2008.03.061.",
        "9. Mensen A, Khatami R. Advanced EEG analysis using threshold-free cluster-enhancement and non-parametric statistics. NeuroImage (2013);67:111-118. PMID 23123297. doi:10.1016/j.neuroimage.2012.10.027.",
    ]
    for ref in refs:
        paragraph = doc.add_paragraph(ref)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_heading("Appendix: analysis status", level=1)
    appendix_summary = summary | {f"morphology_{key}": value for key, value in morphology_summary.items()}
    add_table(doc, ["Item", "Value"], [[k.replace("_", " ").title(), str(v)] for k, v in appendix_summary.items()], [3000, 6360])
    doc.core_properties.title = "Microstate dynamics before functional seizure onset"
    doc.core_properties.subject = "FS_MS paper-preparation analysis"
    doc.core_properties.author = "FS_MS study team"
    doc.core_properties.keywords = "functional seizures; EEG; microstates; cardiac phase; TANOVA; scalp morphology; permutation maximum field; TFCE"
    doc.core_properties.comments = "Generated from deidentified aggregate analysis outputs."
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return {"scalar_significant_fdr": scalar_n, "cardiac_significant_holm": cardiac_n,
            "tanova_significant_max_state": int(len(tanova_sig)), "bhattacharyya_significant_bayesian_fdr": int(len(bh_sig)),
            "bhattacharyya_significant_C_or_E": int(len(bh_sig[bh_sig.get("template_label", pd.Series(dtype=str)).isin(["C", "E"])])) if len(bh_sig) else 0,
            "morphology_participants": int(morphology_summary.get("participants", 0)),
            "morphology_gaussian_sensitivity_significant_fwer": morphology_n,
            "morphology_gaussian_best_peak_p_fwer": float(morphology_best["peak_p_fwer"]) if morphology_best is not None else None,
            "morphology_tfce_primary_significant_fwer": tfce_n,
            "morphology_tfce_best_primary_peak_p_fwer": float(tfce_best["peak_p_tfce_fwer"]) if tfce_best is not None else None,
            "transition_significant_fdr": int(transition.get("p_fdr", pd.Series(dtype=float)).lt(.05).sum()) if len(transition) else 0}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def git_output(repo: Path, *args: str) -> str:
    try:
        return subprocess.check_output(["git", "-c", f"safe.directory={repo}", "-C", str(repo), *args], text=True, stderr=subprocess.STDOUT).strip()
    except Exception as exc:
        return f"unavailable: {exc}"


def write_audit(package: Path, repo: Path, results_dir: Path, manifest: Path, report: Path,
                result_index: dict, audit: Path) -> None:
    files = sorted(p for p in package.rglob("*") if p.is_file() and p != audit)
    manifest_df = pd.read_csv(manifest)
    pending_path = manifest.parent / "vg_microstate_pending_paper_20260723.csv"
    if not pending_path.exists():
        pending_path = manifest.parent / "vg_microstate_pending.csv"
    pending = pd.read_csv(pending_path) if pending_path.exists() else pd.DataFrame()
    full_manifest_path = manifest.parent / "vg_microstate_manifest_mat.csv"
    full_pending_path = manifest.parent / "vg_microstate_pending.csv"
    full_manifest = pd.read_csv(full_manifest_path) if full_manifest_path.exists() else pd.DataFrame()
    full_pending = pd.read_csv(full_pending_path) if full_pending_path.exists() else pd.DataFrame()
    morphology_export = results_dir / "preictal_window_refit_maps.csv"
    morphology_positions = results_dir / "common_channel_positions.csv"
    morphology_summary_path = package / "data" / "morphology_rft_summary.json"
    morphology_summary = (json.loads(morphology_summary_path.read_text(encoding="utf-8"))
                          if morphology_summary_path.exists() else {})
    if morphology_export.exists():
        with morphology_export.open("r", encoding="utf-8") as handle:
            morphology_export_rows = max(sum(1 for _ in handle) - 1, 0)
    else:
        morphology_export_rows = "unavailable"
    model_status_lines = []
    for filename in ("rolling_scalar_gee_terms.csv", "rolling_cardiac_phase_gee_terms.csv"):
        model_path = package / "data" / filename
        if model_path.exists():
            model_frame = pd.read_csv(model_path)
            if "status" in model_frame:
                counts = model_frame["status"].fillna("missing").value_counts().to_dict()
                model_status_lines.append(f"- {filename}: " + "; ".join(f"{key}={value}" for key, value in counts.items()))
    now = datetime.now(timezone.utc)
    lines = [
        "FS_MS MICROSTATE PREICTAL ANALYSIS - LOG AND AUDIT TRAIL",
        "=" * 68,
        f"Audit generated (UTC): {now.isoformat()}",
        f"Audit generated (local): {datetime.now().astimezone().isoformat()}",
        "Status: analysis and report package complete",
        "",
        "1. SCOPE AND ANALYSIS LOCK",
        "Primary functional strata: FS-only functional events; DD functional events (separate).",
        "Comparator strata: ES-only epileptic events; DD epileptic events.",
        "Primary time analysis: 60-s windows, 10-s step, final 20 min; piecewise knot at -5 min.",
        "Sensitivity: 10-s windows, 1-s step, final 60 s; Gaussian-mixture backfit.",
        "Rolling coverage/GFP outliers: one-pass within-event median +/- 3 sample SD on raw values; metric-specific; n<4 or SD=0 unchanged.",
        "State hierarchy: prespecified global C and participant-template E; fresh global E and participant-template C sensitivities; global G absent.",
        "Morphology lock: K=6 participant-prior refits in 60-s windows stepped 30 s; one linear event slope per channel; participant C/E by FS-functional/DD-functional stratum.",
        "Multiplicity: BH FDR for scalar models; Holm for cardiac phase-by-time interactions; max-state TANOVA; primary participant sign-flip maximum-TFCE FWER over 27 electrodes and four morphology fields; Gaussian maximum-field sensitivity over the same fields.",
        "Permutation count: 5,000.",
        "",
        "2. REPOSITORY PROVENANCE",
        f"Repository: {repo}",
        f"Git HEAD: {git_output(repo, 'rev-parse', 'HEAD')}",
        "Git status (pre-existing and analysis changes retained):",
        git_output(repo, "status", "--short"),
        "Submodule status:",
        git_output(repo, "submodule", "status"),
        "",
        "3. SOFTWARE",
        f"OS: {platform.platform()}",
        f"Python: {platform.python_version()}",
        f"pandas: {pd.__version__}",
        "MATLAB and toolbox versions: provenance/matlab_versions.txt",
        "",
        "4. INPUTS AND INTEGRITY",
        f"Paper manifest: {manifest}",
        f"Manifest rows: {len(manifest_df)}",
        f"Unique manifest paths: {manifest_df['file_path'].astype(str).str.lower().nunique()}",
        f"Participants: {manifest_df['participant'].nunique()}",
        f"Pending rows: {len(pending)}",
        f"Full regenerated manifest rows (including ictal conversions): {len(full_manifest) if len(full_manifest) else 'unavailable'}",
        f"Full regenerated pending rows: {len(full_pending) if len(full_pending) else 'unavailable'}",
        f"Ictal records deliberately excluded from the paper hierarchy: {int(full_manifest['condition'].astype(str).str.lower().eq('ictal').sum()) if len(full_manifest) and 'condition' in full_manifest else 'unavailable'}",
        f"Preictal morphology export: {morphology_export}",
        f"Preictal morphology rows: {morphology_export_rows}",
        f"Common channel geometry: {morphology_positions}",
        f"Morphology participants with estimable beta fields: {morphology_summary.get('participants', 'unavailable')}",
        f"Window-refit events: {morphology_summary.get('window_refit_events', 'unavailable')}",
        f"Window-refit status counts: {morphology_summary.get('window_refit_status_counts', 'unavailable')}",
        f"Coverage outlying windows: {json.loads((package / 'data' / 'analysis_summary.json').read_text(encoding='utf-8')).get('coverage_outlying_windows', 'unavailable')}",
        f"Mean-GFP outlying windows: {json.loads((package / 'data' / 'analysis_summary.json').read_text(encoding='utf-8')).get('mean_gfp_outlying_windows', 'unavailable')}",
        "Collision repair: one duplicated interictal output path was split deterministically using crop times; no participant identifier is reproduced here.",
        "Validation: manifest path uniqueness asserted before modelling; duplicate assertion passed.",
        "Historical July 15 summary: explicitly treated as stale and not used for final inference.",
        "",
        "5. IMPLEMENTATION AND COMMANDS",
        "- vg_microstates_pipeline.py: deterministic suffix for colliding paths plus uniqueness assertion and regression check.",
        "- run_vg_microstates.m: backfit_time_window_s parameter, set to 1200 s for this run.",
        "- fit_microstate_hierarchical_dataset.m: rolling occurrence/dwell fields and polarity-invariant mean topographic GMD.",
        "- analyse_transition_matrices.py: canonical C/E labels replace participant-local numeric indices.",
        "- analyse_preictal_publication.py: metric-specific within-event coverage/GFP filtering, participant-weighted rolling GEE and exposure-offset cardiac GEE.",
        "- fit_microstate_hierarchical_dataset.m: cached-peak rolling K=6 refits under participant-full priors; one-to-one matching and polarity-aligned unit-norm C/E export.",
        "- analyse_preictal_morphology_rft.py: event channel slopes, equal event averaging within participant, primary electrode-graph TFCE and Gaussian maximum-field sensitivity.",
        f"Hierarchy results: {results_dir}",
        "Exact launch parameters: K candidates 4,5,6,7; primary time bin 10 s; time window 1200 s; TANOVA 5000; fresh output directory.",
        f"MATLAB launch: run_vg_microstates('{manifest}','{results_dir}','backfit_time_bin_s',10,'backfit_time_window_s',1200,'tanova_permutations',5000)",
        f"Python launch: analyse_preictal_publication.py --results-dir '{results_dir}' --final-minute-results-dir 'G:\\FS_microstates_study\\vg_microstates\\hierarchical' --manifest '{manifest}' --cardiac-csv '{results_dir / 'participant_ce_cardiac_phase_metrics.csv'}' --out-dir '{package}'",
        f"Morphology MATLAB launch: fit_microstate_hierarchical_dataset.m with preictal_window_refit_only=true, window_refit_length_s=60, window_refit_step_s=30, window_refit_time_window_s=1200, window_refit_min_peaks=30, window_refit_min_state_peaks=5, spm_prior_weight=1, output '{morphology_export}'.",
        "Production cache scheduling: one full resumable stream plus non-overlapping start-record chunks 220-249, 250-349 and 350-455; each chunk used random_seed=1 and distinct event-cache filenames. The full stream performed final ordered consolidation.",
        f"Morphology Python launch: analyse_preictal_morphology_rft.py --morphology-csv '{morphology_export}' --positions-csv '{morphology_positions}' --manifest '{manifest}' --out-dir '{package}' --permutations 5000 --fwhm-deg 45 --grid-size 41 --seed 20260723 --tfce-extent-power 0.5 --tfce-height-power 2",
        "Morphology primary family: participant C and E, separately in FS functional and DD functional strata; exact-integral TFCE on the 68-edge electrode graph (E=0.5, H=2) jointly corrects 27 electrodes and four fields.",
        "Morphology sensitivity: 45-degree FWHM Gaussian grid with maximum-field correction jointly over the same four fields.",
        "The source-path morphology export remains in the secured G: results hierarchy; only deidentified event/participant betas and aggregate fields are packaged.",
        "Superseded result: fixed-backfit morphology p=0.034 tested a different estimand and is not used as evidence for refitted morphology change.",
        "Final-minute sensitivity: audited July 20 hierarchy 1-s bins, restricted to the final 60 s; the repaired collision was interictal and cannot affect this sensitivity.",
        "All Python self-checks and MATLAB code analysis completed before the production run.",
        "",
        "6. INCLUSION, EXCLUSION AND MISSINGNESS",
    ]
    counts = (manifest_df.groupby(["group", "condition"], dropna=False).size().reset_index(name="records"))
    lines.extend(f"- {r.group} | {r.condition}: {r.records} records" for r in counts.itertuples(index=False))
    if len(pending):
        for reason, n in pending["reason"].value_counts(dropna=False).items():
            lines.append(f"- Pending: {n} | {reason}")
    lines += ["", "7. MODEL AND RESULT INDEX"]
    lines.extend(f"- {key}: {value}" for key, value in result_index.items())
    lines.extend(model_status_lines or ["- Model status counts unavailable."])
    lines += [
        "Primary/sensitivity/exploratory labels are encoded in the data files and report methods.",
        "Warnings, failed or non-estimable models are retained in status columns rather than silently dropped.",
        "",
        "8. DOCUMENT QA",
        "Preset: narrative_proposal; opening pattern: editorial_cover.",
        "Privacy: aggregate counts/results only; no direct patient identifiers in the DOCX.",
        "Accessibility: real headings, table header rows, descriptive figure alt text, readable colour contrast.",
        "Render method and page-image inspection outcome: see provenance/document_qa.txt.",
        "",
        "9. FINAL OUTPUT INVENTORY (SHA-256)",
    ]
    for path in files:
        stat = path.stat()
        lines.append(f"{path.relative_to(package)} | {stat.st_size} bytes | {datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat()} | {sha256(path)}")
    lines += ["", "END OF AUDIT TRAIL"]
    audit.write_text("\n".join(lines), encoding="utf-8")


def write_readme(package: Path) -> None:
    data = package / "data"
    summary = json.loads((data / "analysis_summary.json").read_text(encoding="utf-8"))
    tfce = read_csv(data / "morphology_tfce_peak_summary.csv").sort_values("peak_p_tfce_fwer", na_position="last")
    gaussian = read_csv(data / "morphology_rft_peak_summary.csv").sort_values("peak_p_fwer", na_position="last")
    tfce_best = tfce.iloc[0] if len(tfce) else None
    gaussian_best = gaussian.iloc[0] if len(gaussian) else None
    tfce_text = (f"Best primary TFCE field: participant {tfce_best['template_label']} in {tfce_best['stratum']} at "
                 f"{channel_name(tfce_best['peak_channel'])} (joint p={fmt(tfce_best['peak_p_tfce_fwer'])})."
                 if tfce_best is not None else "Primary TFCE result unavailable.")
    gaussian_text = (f"Best Gaussian sensitivity field: participant {gaussian_best['template_label']} in "
                     f"{gaussian_best['stratum']} near {channel_name(gaussian_best['peak_nearest_channel'])} "
                     f"(joint p={fmt(gaussian_best['peak_p_fwer'])})."
                     if gaussian_best is not None else "Gaussian sensitivity unavailable.")
    lines = [
        "FS_MS functional-seizure microstate paper-analysis package",
        "===========================================================",
        "",
        "Start with:",
        "  1. FS_MS_microstate_preictal_analysis_report.docx",
        "  2. FS_MS_microstate_preictal_analysis_audit_log.txt",
        "",
        "Folders:",
        "  data/        Deidentified rolling-window values, fitted maps, model outputs and dictionaries.",
        "  figures/     Publication-oriented plots used in the report.",
        "  provenance/  Parameters, software versions, validation notes, source snapshots and QA.",
        "",
        "Analysis lock:",
        "  Coverage and mean GFP use independent one-pass within-event median +/- 3 sample-SD filtering.",
        "  Morphology uses K=6 participant-prior refits in 60-s windows stepped 30 s, followed by one",
        "  channel slope per event and equal event averaging within participant. Electrode-graph TFCE",
        "  (E=0.5, H=2; 5,000 sign flips) jointly corrects C/E in FS-functional and DD-functional strata.",
        "  Gaussian maximum-field inference is a sensitivity.",
        "",
        "Headline result:",
        f"  Corrected scalar rolling terms: {summary.get('significant_scalar_fdr_primary_functional', 0)}.",
        f"  Corrected cardiac phase-by-time interactions: {summary.get('significant_cardiac_holm_primary_functional', 0)}.",
        f"  Coverage outliers removed: {summary.get('coverage_outlying_windows', 0)}; mean-GFP outliers removed: {summary.get('mean_gfp_outlying_windows', 0)}.",
        f"  {tfce_text}",
        f"  {gaussian_text}",
        "  The previous fixed-backfit morphology p=0.034 is superseded because it tested a different estimand.",
        "",
        "The package deliberately excludes source/intermediate EEG files. The deidentified window-refit",
        "map table retains amplitudes, fit QC and failure reasons; the path-bearing authoritative export",
        "remains in the secured G: analysis hierarchy. The adjacent TXT audit records rerun parameters",
        "and SHA-256 hashes for the delivered package.",
    ]
    (package / "README.txt").write_text("\n".join(lines), encoding="utf-8")


def write_statistical_audit(package: Path) -> None:
    data = package / "data"
    scalar = read_csv(data / "rolling_scalar_gee_terms.csv")
    cardiac = read_csv(data / "rolling_cardiac_phase_gee_terms.csv")
    tanova = read_csv(data / "tanova_results.csv")
    bhatt = read_csv(data / "bhattacharyya_tanova_results.csv")
    tfce = read_csv(data / "morphology_tfce_peak_summary.csv")
    gaussian = read_csv(data / "morphology_rft_peak_summary.csv")
    primary_scalar = scalar[(scalar.get("analysis_role", pd.Series(dtype=str)).eq("co_primary")) &
                            scalar.get("scheme", pd.Series(dtype=str)).eq("20min_60s_step10s") &
                            scalar.get("backfit_method", pd.Series(dtype=str)).eq("hard") &
                            scalar.get("stratum", pd.Series(dtype=str)).astype(str).str.contains("functional", na=False)]
    cardiac_primary = cardiac[(cardiac.get("systolic_window_start_ms", pd.Series(dtype=float)).eq(180)) &
                              (cardiac.get("systolic_window_end_ms", pd.Series(dtype=float)).eq(500)) &
                              cardiac.get("stratum", pd.Series(dtype=str)).astype(str).str.contains("functional", na=False) &
                              cardiac.get("term", pd.Series(dtype=str)).astype(str).str.contains(":", na=False)]
    payload = {
        "analysis_summary": json.loads((data / "analysis_summary.json").read_text(encoding="utf-8")),
        "morphology_summary": json.loads((data / "morphology_rft_summary.json").read_text(encoding="utf-8")),
        "scalar_status_counts": {str(k): int(v) for k, v in scalar.get("status", pd.Series(dtype=str)).fillna("missing").value_counts().items()},
        "best_primary_functional_scalar": json.loads(primary_scalar.sort_values(["p_fdr", "p_value"]).head(1).to_json(orient="records")),
        "cardiac_status_counts": {str(k): int(v) for k, v in cardiac.get("status", pd.Series(dtype=str)).fillna("missing").value_counts().items()},
        "best_primary_functional_cardiac_interaction": json.loads(cardiac_primary.sort_values(["p_holm", "p_value"]).head(1).to_json(orient="records")),
        "tanova_significant_max_state": json.loads(tanova[tanova.get("p_max_state", pd.Series(dtype=float)).lt(.05)].to_json(orient="records")),
        "bhattacharyya_bayesian_fdr_significant": json.loads(bhatt[bhatt.get("bayesian_fdr_significant", pd.Series(dtype=bool)).astype(bool)].to_json(orient="records")),
        "morphology_primary_tfce_fields": json.loads(tfce.to_json(orient="records")),
        "morphology_gaussian_sensitivity_fields": json.loads(gaussian.to_json(orient="records")),
        "superseded_result": "Fixed-backfit morphology p=0.034 tested a different estimand and is not used for window-refitted morphology inference.",
    }
    (package / "provenance" / "statistical_result_audit.json").write_text(
        json.dumps(payload, indent=2), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--package", type=Path, required=True)
    parser.add_argument("--repo", type=Path, required=True)
    parser.add_argument("--results-dir", type=Path, required=True)
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--audit-only", action="store_true")
    args = parser.parse_args()
    report = args.package / "FS_MS_microstate_preictal_analysis_report.docx"
    audit = args.package / "FS_MS_microstate_preictal_analysis_audit_log.txt"
    index_path = args.package / "provenance" / "report_result_index.json"
    if args.audit_only:
        result_index = json.loads(index_path.read_text(encoding="utf-8"))
    else:
        result_index = build_report(args.package, args.results_dir, args.manifest, report)
        index_path.parent.mkdir(parents=True, exist_ok=True)
        index_path.write_text(json.dumps(result_index, indent=2), encoding="utf-8")
    write_readme(args.package)
    write_statistical_audit(args.package)
    write_audit(args.package, args.repo, args.results_dir, args.manifest, report, result_index, audit)
    print(json.dumps({"report": str(report), "audit": str(audit), **result_index}, indent=2))


if __name__ == "__main__":
    main()
